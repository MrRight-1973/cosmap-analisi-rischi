"""Confronto tra la prima estrazione (Libreria_analisi_rischi_Cosmap.xlsx) e la nuova libreria.

Uso:  python dati/confronta_librerie.py
"""

import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).parent))
import genera_libreria_nuova as nuova  # noqa: E402

VECCHIA = Path(__file__).with_name("Libreria_analisi_rischi_Cosmap.xlsx")
USCITA = Path(__file__).with_name("Confronto_librerie.xlsx")
RE_NORMA = re.compile(r"\b(EN ISO/TR|ISO/TR|EN ISO|EN IEC|IEC/IEEE|IEC/TS|EN|ISO|IEC)\s+(\d+(?:-\d+)*)")


def intero(v):
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return None


def leggi_vecchia():
    wb = openpyxl.load_workbook(VECCHIA, data_only=True)
    ress_dir_reg = {}
    for r in wb["RESS-Regolamento"].iter_rows(min_row=2, values_only=True):
        if r[4]:
            ress_dir_reg[str(r[0]).strip()] = str(r[4]).strip()
    schede = []
    for r in wb["Libreria"].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        si = tuple(intero(v) for v in r[10:14])
        sf = tuple(intero(v) for v in r[18:22])
        ress = str(r[1]).strip()
        schede.append(dict(
            codice=r[0], ress=ress_dir_reg.get(ress, ress), modulo=str(r[5]).strip().replace("Zona smerigliatura – generale", "Zona smerigliatura - generale"),
            si=si if None not in si else None, sf=sf if None not in sf else None,
            misure=str(r[16] or ""), norme={f"{a} {b}" for a, b in RE_NORMA.findall(str(r[24] or ""))},
        ))
    return schede


def stima_txt(st):
    if not st:
        return "–"
    cl, es = nuova.esito(st)
    return f"Se{st[0]} Fr{st[1]} Pr{st[2]} Av{st[3]} → Cl {cl} {es}"


def esito_finale(scheda):
    return nuova.esito(scheda["sf"])[1] if scheda["sf"] else None


def main():
    vecchie = leggi_vecchia()
    nuove = [dict(s, codice=f"NL-{i:03d}") for i, s in enumerate(nuova.SCHEDE, start=1)]
    wb = openpyxl.Workbook()
    grassetto = Font(bold=True)
    fondo = PatternFill("solid", fgColor="DDE4EE")
    evid = PatternFill("solid", fgColor="FFF2CC")

    def foglio(titolo, colonne, larghezze, prima=False):
        ws = wb.active if prima else wb.create_sheet()
        ws.title = titolo
        ws.append(colonne)
        for c in ws[1]:
            c.font, c.fill = grassetto, fondo
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(larghezze, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return ws

    # Riepilogo
    ws = foglio("Riepilogo", ["Voce", "Libreria attuale", "Nuova libreria", "Commento"], [42, 18, 18, 80], prima=True)
    tipi = Counter(t for s in nuove for t, _, _ in s["mis"])
    misure_n = sum(len(s["mis"]) for s in nuove)
    righe = [
        ("Schede", len(vecchie), len(nuove), "La nuova ha una scheda per pericolo e zona: più schede per lo stesso requisito."),
        ("Schede con stima", sum(1 for s in vecchie if s["si"]), sum(1 for s in nuove if s["si"]), ""),
        ("Esito finale OK", sum(1 for s in vecchie if esito_finale(s) == "OK"),
         sum(1 for s in nuove if esito_finale(s) == "OK"),
         "Nella nuova le schede con Se 3-4 restano spesso 'suggerite', con il rischio residuo dichiarato."),
        ("Esito finale non verde", sum(1 for s in vecchie if s["sf"] and esito_finale(s) != "OK"),
         sum(1 for s in nuove if s["sf"] and esito_finale(s) != "OK"), ""),
        ("Misure", sum(1 for s in vecchie if s["misure"].strip()), misure_n,
         "Attuale: un unico testo per scheda. Nuova: una misura per riga."),
        ("Misure classificate per tipo", 0, misure_n,
         f"Progettazione {tipi['PROG']}, protezione {tipi['PROT']}, informazioni {tipi['INFO']}."),
        ("Misure con norma di riferimento", 0, sum(1 for s in nuove for _, _, n in s["mis"] if n), ""),
        ("Requisiti RESS con almeno una scheda", len({s["ress"] for s in vecchie}), len({s["ress"] for s in nuove}), ""),
        ("Norme citate", len(set().union(*(s["norme"] for s in vecchie))), len(nuova.NORME), ""),
    ]
    for r in righe:
        ws.append(list(r))

    # Requisiti
    ws = foglio("Requisiti", ["RESS", "Titolo", "Schede attuali", "Schede nuove", "Differenza"], [10, 55, 14, 14, 50])
    per_v = Counter(s["ress"] for s in vecchie)
    per_n = Counter(s["ress"] for s in nuove)
    codici_nuovi = {r[2] for r in nuova.RESS}
    for _, titolo, codice, _, _ in nuova.RESS:
        v, n = per_v.get(codice, 0), per_n.get(codice, 0)
        nota = ""
        if v and not n:
            nota = "Trattato solo nella libreria attuale"
        elif n and not v:
            nota = "Trattato solo nella nuova libreria"
        ws.append([codice, titolo, v, n, nota])
        if nota:
            for c in ws[ws.max_row]:
                c.fill = evid
    for codice in sorted(set(per_v) - codici_nuovi):
        ws.append([codice, "Fuori dalla parte 1 dell'Allegato III", per_v[codice], 0, "Da verificare"])

    # Moduli
    ws = foglio("Moduli", ["Modulo", "Schede attuali", "Schede nuove", "RESS solo attuali", "RESS solo nuove"], [50, 14, 14, 45, 45])
    for nome, _, _ in nuova.MODULI:
        rv = {s["ress"] for s in vecchie if s["modulo"] == nome}
        rn = {s["ress"] for s in nuove if s["modulo"] == nome}
        ws.append([nome, sum(1 for s in vecchie if s["modulo"] == nome), sum(1 for s in nuove if s["modulo"] == nome),
                   ", ".join(sorted(rv - rn)), ", ".join(sorted(rn - rv))])

    # Stime a confronto per modulo e requisito
    ws = foglio("Stime a confronto", ["Modulo", "RESS", "Scheda attuale", "Stima iniziale attuale", "Stima finale attuale",
                                      "Scheda nuova", "Stima iniziale nuova", "Stima finale nuova", "Differenza"],
                [34, 9, 10, 32, 32, 10, 32, 32, 45])
    gruppi = defaultdict(lambda: ([], []))
    for s in vecchie:
        gruppi[(s["modulo"], s["ress"])][0].append(s)
    for s in nuove:
        gruppi[(s["modulo"], s["ress"])][1].append(s)
    ordine_mod = {nome: i for i, (nome, _, _) in enumerate(nuova.MODULI)}
    for (modulo, ress), (lv, ln) in sorted(gruppi.items(), key=lambda kv: (ordine_mod.get(kv[0][0], 99), kv[0][1])):
        for i in range(max(len(lv), len(ln))):
            v = lv[i] if i < len(lv) else None
            n = ln[i] if i < len(ln) else None
            diff = ""
            if v and n:
                if v["si"] and n["si"] and v["si"][0] != n["si"][0]:
                    diff = f"Gravità diversa (Se {v['si'][0]} → {n['si'][0]})"
                if esito_finale(v) != esito_finale(n):
                    diff = (diff + "; " if diff else "") + f"Esito finale {esito_finale(v) or '–'} → {esito_finale(n) or '–'}"
            elif v:
                diff = "Solo nella libreria attuale"
            else:
                diff = "Solo nella nuova libreria"
            ws.append([modulo, ress, v["codice"] if v else "", stima_txt(v["si"]) if v else "", stima_txt(v["sf"]) if v else "",
                       n["codice"] if n else "", stima_txt(n["si"]) if n else "", stima_txt(n["sf"]) if n else "", diff])
            if diff:
                ws.cell(ws.max_row, 9).fill = evid

    # Norme
    ws = foglio("Norme", ["Norma", "Libreria attuale", "Nuova libreria"], [24, 16, 16])
    nv = set().union(*(s["norme"] for s in vecchie))
    nn = {n[0] for n in nuova.NORME}
    for codice in sorted(nv | nn):
        ws.append([codice, "sì" if codice in nv else "", "sì" if codice in nn else ""])
        if (codice in nv) != (codice in nn):
            for c in ws[ws.max_row]:
                c.fill = evid

    wb.save(USCITA)
    print(f"Creato {USCITA.name}")
    for r in righe:
        print(r[:3])
    print("RESS solo attuale:", sorted(c for c in per_v if c in codici_nuovi and not per_n.get(c)))
    print("RESS solo nuova:", sorted(c for c in per_n if not per_v.get(c)))
    print("Norme solo attuale:", sorted(nv - nn))
    print("Norme solo nuova:", sorted(nn - nv))
    print("Fuori parte 1:", sorted(set(per_v) - codici_nuovi))


if __name__ == "__main__":
    main()
