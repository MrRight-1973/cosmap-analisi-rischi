import io
from pathlib import Path

from django.conf import settings
from django.contrib.auth.models import Group, User
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse

from . import servizi
from .models import (
    Cliente,
    Commessa,
    Esito,
    Macchina,
    MetodoStima,
    Modulo,
    RegistroModifica,
    Revisione,
    RevisioneBloccata,
    RiferimentoNormativo,
    SchedaAnalisi,
    SchedaModello,
)
from .registro import utente_attivo

LIBRERIA = Path(settings.BASE_DIR) / "dati" / "Libreria_analisi_rischi_Cosmap.xlsx"


def crea_utente(nome, *ruoli):
    utente = User.objects.create_user(nome, password="prova-prova-123")
    for ruolo in ruoli:
        utente.groups.add(Group.objects.get_or_create(name=ruolo)[0])
    return utente


class BaseConLibreria(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("importa_libreria", str(LIBRERIA), stdout=io.StringIO())
        cls.compilatore = crea_utente("mario", servizi.COMPILATORE)
        cls.verificatore = crea_utente("lucia", servizi.VERIFICATORE)
        cls.approvatore = crea_utente("titolare", servizi.APPROVATORE)
        cls.riferimento = RiferimentoNormativo.objects.get()
        cls.cliente = Cliente.objects.create(ragione_sociale="Kohler")

    def nuova_macchina(self, numero="25-20", *moduli):
        commessa = Commessa.objects.create(
            numero=numero, cliente=self.cliente, anno=2020, riferimento=self.riferimento
        )
        macchina = Macchina.objects.create(commessa=commessa, denominazione="Impianto di prova", matricola="X1")
        macchina.moduli.set(servizi.moduli_proposti())
        if moduli:
            macchina.moduli.add(*Modulo.objects.filter(nome__startswith=moduli[0]))
        return macchina


class ImportTest(BaseConLibreria):
    def test_libreria_importata(self):
        self.assertEqual(SchedaModello.objects.count(), 78)
        self.assertEqual(self.riferimento.requisitoress_set.count(), 63)

    def test_matrice_come_foglio_metodo(self):
        metodo = MetodoStima.corrente()
        attesi = {
            4: ["SUGGERITE", "RICHIESTE", "RICHIESTE", "RICHIESTE", "RICHIESTE"],
            3: ["OK", "SUGGERITE", "RICHIESTE", "RICHIESTE", "RICHIESTE"],
            2: ["OK", "OK", "SUGGERITE", "RICHIESTE", "RICHIESTE"],
            1: ["OK", "OK", "OK", "SUGGERITE", "RICHIESTE"],
        }
        for se, esiti in attesi.items():
            self.assertEqual([metodo.esito(se, cl) for cl in (4, 6, 9, 12, 15)], esiti)

    def test_reimport_non_duplica(self):
        call_command("importa_libreria", str(LIBRERIA), stdout=io.StringIO())
        self.assertEqual(SchedaModello.objects.count(), 78)


class LibreriaNuovaTest(TestCase):
    def test_import_con_misure_classificate(self):
        import io

        from .models import MisuraModello, Norma, TipoMisura

        call_command("importa_libreria", str(LIBRERIA.parent / "Libreria_nuova_Cosmap.xlsx"), stdout=io.StringIO())
        self.assertEqual(Modulo.objects.count(), 14)
        self.assertGreater(SchedaModello.objects.count(), 80)
        self.assertFalse(MisuraModello.objects.filter(tipo=TipoMisura.DA_CLASSIFICARE).exists())
        self.assertTrue(Norma.objects.get(codice="EN ISO 12100").armonizzata)
        self.assertFalse(Norma.objects.get(codice="EN IEC 62443-3-3").armonizzata)
        self.assertEqual(Norma.objects.get(codice="EN ISO 10218-2").tipo, Norma.Tipo.C)
        self.assertTrue(MisuraModello.objects.filter(scheda__codice="NL-031", norma__codice="EN ISO 13855").exists())
        metodo = MetodoStima.corrente()
        for scheda in SchedaModello.objects.all():
            esito = metodo.esito(scheda.se_finale, scheda.cl_finale)
            self.assertNotEqual(esito, Esito.RICHIESTE, scheda.codice)
            if esito and esito != Esito.OK:
                self.assertTrue(scheda.testo_istruzioni, scheda.codice)


class CreazioneAnalisiTest(BaseConLibreria):
    def test_solo_moduli_scelti(self):
        sempre = SchedaModello.objects.filter(modulo__sempre_attivo=True).count()
        analisi = servizi.crea_analisi_da_libreria(self.nuova_macchina(), self.compilatore)
        self.assertEqual(analisi.revisione_corrente.schede.count(), sempre)

        tavola = "Tavola rotante"
        con_tavola = SchedaModello.objects.filter(modulo__nome__startswith=tavola).count()
        self.assertTrue(con_tavola)
        altra = servizi.crea_analisi_da_libreria(self.nuova_macchina("26-01", tavola), self.compilatore)
        self.assertEqual(altra.revisione_corrente.schede.count(), sempre + con_tavola)

    def test_copia_indipendente_dalla_libreria(self):
        analisi = servizi.crea_analisi_da_libreria(self.nuova_macchina(), self.compilatore)
        scheda = analisi.revisione_corrente.schede.filter(origine__isnull=False).first()
        modello = scheda.origine
        modello.testo_istruzioni = "Testo cambiato in libreria"
        modello.save()
        scheda.refresh_from_db()
        self.assertNotEqual(scheda.testo_istruzioni, "Testo cambiato in libreria")
        self.assertEqual(scheda.misure.count(), modello.misure.count())

    def test_solo_compilatore_crea(self):
        with self.assertRaises(PermissionDenied):
            servizi.crea_analisi_da_libreria(self.nuova_macchina(), self.verificatore)


class FlussoTest(BaseConLibreria):
    def setUp(self):
        self.analisi = servizi.crea_analisi_da_libreria(self.nuova_macchina(), self.compilatore)
        self.rev = self.analisi.revisione_corrente

    def test_esito_calcolato(self):
        scheda = self.rev.schede.first()
        scheda.se_finale, scheda.fr_finale, scheda.pr_finale, scheda.av_finale = 3, 2, 1, 1
        self.assertEqual(scheda.cl_finale, 4)
        self.assertEqual(scheda.esito_finale, Esito.OK)

    def test_scartare_richiede_motivazione(self):
        scheda = self.rev.schede.first()
        with self.assertRaises(ValidationError):
            servizi.decidi_scheda(scheda, self.compilatore, SchedaAnalisi.Decisione.SCARTATA)
        servizi.decidi_scheda(scheda, self.compilatore, SchedaAnalisi.Decisione.SCARTATA, "Non presente")
        self.assertEqual(scheda.decisa_da, self.compilatore)

    def test_flusso_completo_e_blocco(self):
        servizi.invia_in_verifica(self.rev, self.compilatore)
        with self.assertRaises(RevisioneBloccata):
            self.rev.schede.first().save()
        with self.assertRaises(ValidationError):
            servizi.approva(self.rev, self.approvatore)  # non ancora verificata
        servizi.segna_verificata(self.rev, self.verificatore)
        servizi.approva(self.rev, self.approvatore)
        self.assertEqual(self.rev.stato, Revisione.Stato.APPROVATA)

        scheda = self.rev.schede.first()
        with self.assertRaises(RevisioneBloccata):
            scheda.pericoli.clear()
        with self.assertRaises(RevisioneBloccata):
            scheda.delete()

    def test_chi_compila_non_approva(self):
        tuttofare = crea_utente("tutto", servizi.COMPILATORE, servizi.VERIFICATORE, servizi.APPROVATORE)
        analisi = servizi.crea_analisi_da_libreria(self.nuova_macchina("26-02"), tuttofare)
        rev = analisi.revisione_corrente
        servizi.invia_in_verifica(rev, tuttofare)
        servizi.segna_verificata(rev, tuttofare)
        with self.assertRaises(PermissionDenied):
            servizi.approva(rev, tuttofare)

    def test_rimanda_in_bozza(self):
        servizi.invia_in_verifica(self.rev, self.compilatore)
        with self.assertRaises(ValidationError):
            servizi.rimanda_in_bozza(self.rev, self.verificatore, "")
        servizi.rimanda_in_bozza(self.rev, self.verificatore, "Manca la stima finale")
        self.assertTrue(self.rev.modificabile)

    def test_nuova_revisione_copia_e_sostituisce(self):
        scheda = self.rev.schede.first()
        servizi.decidi_scheda(scheda, self.compilatore, SchedaAnalisi.Decisione.SCARTATA, "Non presente")
        servizi.invia_in_verifica(self.rev, self.compilatore)
        servizi.segna_verificata(self.rev, self.verificatore)
        servizi.approva(self.rev, self.approvatore)

        with self.assertRaises(ValidationError):
            servizi.nuova_revisione(self.analisi, self.compilatore, "")
        rev1 = servizi.nuova_revisione(self.analisi, self.compilatore, "Aggiunto robot")
        self.assertEqual(rev1.numero, 1)
        self.assertEqual(rev1.schede.count(), self.rev.schede.count())
        self.assertEqual(rev1.schede.filter(decisione=SchedaAnalisi.Decisione.SCARTATA).count(), 1)
        self.assertEqual(rev1.applicabilita.count(), self.rev.applicabilita.count())

        servizi.invia_in_verifica(rev1, self.compilatore)
        servizi.segna_verificata(rev1, self.verificatore)
        servizi.approva(rev1, self.approvatore)
        self.rev.refresh_from_db()
        self.assertEqual(self.rev.stato, Revisione.Stato.SOSTITUITA)

    def test_controlli_completezza(self):
        tipi = {a.tipo for a in servizi.controlli(self.rev)}
        self.assertIn("decisione", tipi)
        self.assertIn("requisito", tipi)  # es. 1.1.9, requisito nuovo senza schede

    def test_copia_da_altra_analisi(self):
        servizi.invia_in_verifica(self.rev, self.compilatore)
        servizi.segna_verificata(self.rev, self.verificatore)
        servizi.approva(self.rev, self.approvatore)
        copia = servizi.crea_analisi_da_copia(self.nuova_macchina("26-03"), self.rev, self.compilatore)
        rev = copia.revisione_corrente
        self.assertEqual(rev.schede.count(), self.rev.schede.count())
        self.assertFalse(rev.schede.exclude(decisione=SchedaAnalisi.Decisione.PROPOSTA).exists())

    def test_registro_modifiche(self):
        scheda = self.rev.schede.first()
        with utente_attivo(self.compilatore):
            scheda.se_iniziale = 4
            scheda.save()
        voce = RegistroModifica.objects.filter(oggetto_id=str(scheda.pk), azione="MODIFICA").first()
        self.assertEqual(voce.utente, self.compilatore)
        self.assertIn("se_iniziale", voce.modifiche)
        with self.assertRaises(ValidationError):
            voce.delete()


class PagineTest(BaseConLibreria):
    def test_percorso_web(self):
        self.client.force_login(self.compilatore)
        risposta = self.client.post(
            reverse("nuova_commessa"),
            {
                "numero": "27-01",
                "nuovo_cliente": "Cliente prova",
                "anno": 2027,
                "riferimento": self.riferimento.pk,
                "denominazione": "Tavola rotante",
                "tipo": "MACCHINA",
                "origine": "LIBRERIA",
                "moduli": [m.pk for m in servizi.moduli_proposti()],
            },
        )
        analisi = Commessa.objects.get(numero="27-01").macchine.get().analisi
        self.assertRedirects(risposta, reverse("analisi", args=[analisi.pk]))
        self.assertContains(self.client.get(reverse("analisi", args=[analisi.pk])), "Invia in verifica")
        self.assertContains(self.client.get(reverse("elenco_commesse")), "27-01")

        scheda = analisi.revisione_corrente.schede.first()
        self.assertEqual(self.client.get(reverse("scheda", args=[scheda.pk])).status_code, 200)
        self.client.post(reverse("conferma_scheda", args=[scheda.pk]))
        scheda.refresh_from_db()
        self.assertEqual(scheda.decisione, SchedaAnalisi.Decisione.CONFERMATA)
        self.assertEqual(
            self.client.get(reverse("applicabilita", args=[analisi.revisione_corrente.pk])).status_code, 200
        )
        self.assertEqual(self.client.get(reverse("registro")).status_code, 200)

    def test_modifica_scheda_da_form(self):
        self.client.force_login(self.compilatore)
        analisi = servizi.crea_analisi_da_libreria(self.nuova_macchina(), self.compilatore)
        scheda = analisi.revisione_corrente.schede.filter(misure__isnull=False).first()
        pagina = self.client.get(reverse("scheda", args=[scheda.pk]))
        dati = {}
        form = pagina.context["form"]
        for nome, campo in form.fields.items():
            valore = form.initial.get(nome)
            if hasattr(valore, "__iter__") and not isinstance(valore, str):
                dati[nome] = [getattr(v, "pk", v) for v in valore]
            elif valore is not None:
                dati[nome] = getattr(valore, "pk", valore)
        misure = pagina.context["misure"]
        dati.update({f"misure-{k}": v for k, v in misure.management_form.initial.items()})
        dati["misure-TOTAL_FORMS"] = misure.initial_form_count()
        for i, mf in enumerate(misure.forms):
            if mf.instance.pk:
                dati.update({f"misure-{i}-id": mf.instance.pk, f"misure-{i}-ordine": 0, f"misure-{i}-tipo": "PROT", f"misure-{i}-testo": mf.instance.testo})
        dati["se_finale"] = 2
        risposta = self.client.post(reverse("scheda", args=[scheda.pk]), dati)
        self.assertEqual(risposta.status_code, 302, risposta.context and (risposta.context["form"].errors, risposta.context["misure"].errors, risposta.context["misure"].non_form_errors()))
        scheda.refresh_from_db()
        self.assertEqual(scheda.se_finale, 2)
        self.assertEqual(scheda.decisione, SchedaAnalisi.Decisione.MODIFICATA)
        self.assertEqual(scheda.misure.first().tipo, "PROT")


class DocumentiTest(BaseConLibreria):
    def setUp(self):
        import tempfile

        from django.test import override_settings

        self.media = tempfile.TemporaryDirectory()
        self.impostazioni = override_settings(MEDIA_ROOT=self.media.name)
        self.impostazioni.enable()
        macchina = self.nuova_macchina()
        macchina.modello = "TR 5/4 CNC"
        macchina.funzione = "Lucidatura automatica di rubinetteria"
        macchina.save()
        from .models import LegislazioneUE

        macchina.altre_legislazioni.set(LegislazioneUE.objects.filter(predefinita=True))
        self.analisi = servizi.crea_analisi_da_libreria(macchina, self.compilatore)
        self.rev = self.analisi.revisione_corrente

    def tearDown(self):
        self.impostazioni.disable()
        self.media.cleanup()

    @staticmethod
    def leggi(contenuto):
        import io

        import docx

        d = docx.Document(io.BytesIO(contenuto))
        parti = [p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]
        parti += [p.text for s in d.sections for p in s.header.paragraphs]
        return "\n".join(parti)

    def approva(self):
        servizi.invia_in_verifica(self.rev, self.compilatore)
        servizi.segna_verificata(self.rev, self.verificatore)
        servizi.approva(self.rev, self.approvatore)

    def test_dichiarazione_bozza_e_definitiva(self):
        from . import documenti

        bozza = self.leggi(documenti.dichiarazione(self.rev))
        self.assertIn("BOZZA", bozza)
        self.assertIn("Regolamento (UE) 2023/1230", bozza)
        self.assertIn("C.O.S.M.A.P. s.r.l.", bozza)
        self.assertIn("2014/30/UE", bozza)
        self.assertIn("Lucidatura automatica di rubinetteria", bozza)
        self.assertIn("EN ISO 12100", bozza)

        self.approva()
        definitiva = self.leggi(documenti.dichiarazione(self.rev, "en"))
        self.assertNotIn("DRAFT", definitiva)
        self.assertIn("EU Declaration of Conformity", definitiva)
        self.assertIn("Saccolongo", definitiva)

    def test_valutazione_e_residui(self):
        from . import documenti

        scheda = self.rev.schede.exclude(testo_istruzioni="").first()
        servizi.decidi_scheda(scheda, self.compilatore, SchedaAnalisi.Decisione.SCARTATA, "Non presente sulla macchina")
        valutazione = self.leggi(documenti.valutazione(self.rev))
        self.assertIn("Metodo di stima", valutazione)
        self.assertIn("Non presente sulla macchina", valutazione)
        self.assertIn("Misure richieste", valutazione)
        residui = self.leggi(documenti.rischi_residui(self.rev))
        self.assertIn("Rischi residui", residui)
        altra = self.rev.schede.exclude(pk=scheda.pk).exclude(testo_istruzioni="").first()
        self.assertIn(altra.testo_istruzioni[:40], residui)

    def test_stime_con_descrizione(self):
        from . import documenti

        scheda = self.rev.schede.exclude(se_iniziale=None).first()
        self.client.force_login(self.compilatore)
        pagina = self.client.get(reverse("scheda", args=[scheda.pk]))
        self.assertContains(pagina, "4 – Morte")
        self.assertContains(pagina, "3 – Possibile")
        valutazione = self.leggi(documenti.valutazione(self.rev))
        self.assertIn("Valori dei fattori", valutazione)
        se = self.rev.metodo.descrizioni()[("Se", scheda.se_iniziale)]
        self.assertIn(f"{scheda.se_iniziale} – {se}", valutazione)

    def test_genera_da_pagina_e_archivia(self):
        from .models import DocumentoGenerato

        self.client.force_login(self.compilatore)
        risposta = self.client.post(reverse("genera_documento", args=[self.rev.pk, "dichiarazione"]), {"lingua": "en"})
        self.assertEqual(risposta.status_code, 200)
        self.assertIn("_BOZZA_en.docx", risposta["Content-Disposition"])
        risposta.close()  # su Windows un file aperto non si può cancellare
        documento = DocumentoGenerato.objects.get()
        self.assertFalse(documento.definitivo)
        self.assertContains(self.client.get(reverse("analisi", args=[self.analisi.pk])), "Scarica")
        scaricato = self.client.get(reverse("scarica_documento", args=[documento.pk]))
        self.assertEqual(scaricato.status_code, 200)
        scaricato.close()

        self.approva()
        self.client.post(reverse("genera_documento", args=[self.rev.pk, "valutazione"])).close()
        self.assertTrue(DocumentoGenerato.objects.get(tipo="VALUTAZIONE").definitivo)


class ModuliDellaMacchinaTest(BaseConLibreria):
    def setUp(self):
        self.macchina = self.nuova_macchina()
        self.analisi = servizi.crea_analisi_da_libreria(self.macchina, self.compilatore)
        self.rev = self.analisi.revisione_corrente
        self.tavola = Modulo.objects.filter(nome__startswith="Tavola rotante").first()

    def test_aggiunge_e_toglie_moduli(self):
        moduli = list(self.macchina.moduli.all())
        aggiunte, tolte, rimaste = servizi.cambia_moduli(self.macchina, moduli + [self.tavola], self.compilatore)
        self.assertEqual(aggiunte, SchedaModello.objects.filter(modulo=self.tavola).count())
        self.assertEqual((tolte, rimaste), (0, 0))

        decisa = self.rev.schede.filter(modulo=self.tavola).first()
        servizi.decidi_scheda(decisa, self.compilatore, SchedaAnalisi.Decisione.CONFERMATA, "")
        aggiunte, tolte, rimaste = servizi.cambia_moduli(self.macchina, moduli, self.compilatore)
        self.assertEqual((aggiunte, rimaste), (0, 1))
        self.assertEqual(self.rev.schede.filter(modulo=self.tavola).count(), 1)
        self.assertNotIn(self.tavola, self.macchina.moduli.all())

    def test_non_si_cambiano_con_revisione_approvata(self):
        servizi.invia_in_verifica(self.rev, self.compilatore)
        servizi.segna_verificata(self.rev, self.verificatore)
        servizi.approva(self.rev, self.approvatore)
        with self.assertRaises(ValidationError):
            servizi.cambia_moduli(self.macchina, [self.tavola], self.compilatore)

    def test_dalla_pagina_della_macchina(self):
        self.client.force_login(self.compilatore)
        url = reverse("macchina", args=[self.macchina.pk])
        self.assertContains(self.client.get(url), "Moduli della libreria")
        risposta = self.client.post(
            url,
            {
                "denominazione": self.macchina.denominazione,
                "tipo": self.macchina.tipo,
                "moduli": [m.pk for m in self.macchina.moduli.all()] + [self.tavola.pk],
            },
        )
        self.assertRedirects(risposta, reverse("analisi", args=[self.analisi.pk]))
        self.assertIn(self.tavola, self.macchina.moduli.all())
        self.assertTrue(self.rev.schede.filter(modulo=self.tavola).exists())


class AmministrazioneLibreriaTest(BaseConLibreria):
    def test_schede_si_scelgono_dal_modulo(self):


        admin_utente = User.objects.create_superuser("capo", password="prova-prova-123")
        self.client.force_login(admin_utente)
        robot = Modulo.objects.get(nome__startswith="Gruppo di smerigliatura")
        altra = SchedaModello.objects.exclude(modulo=robot).first()
        proprie = list(robot.schede.values_list("pk", flat=True))
        url = reverse("admin:rischi_modulo_change", args=[robot.pk])
        dati = {
            "nome": robot.nome,
            "descrizione": robot.descrizione,
            "condizione": robot.condizione,
            "attivo": "on",
            "ordine": robot.ordine,
        }

        self.assertContains(self.client.get(url), "Schede del modulo")
        risposta = self.client.post(url, {**dati, "schede": proprie + [altra.pk]})
        self.assertEqual(risposta.status_code, 302)
        altra.refresh_from_db()
        self.assertEqual(altra.modulo, robot)

        risposta = self.client.post(url, {**dati, "schede": proprie})
        self.assertContains(risposta, "non può restare senza modulo")
        altra.refresh_from_db()
        self.assertEqual(altra.modulo, robot)


class EliminaCommessaTest(BaseConLibreria):
    def setUp(self):
        self.analisi = servizi.crea_analisi_da_libreria(self.nuova_macchina("26-99"), self.compilatore)
        self.commessa = self.analisi.macchina.commessa
        self.url = reverse("elimina_commessa", args=[self.commessa.pk])

    def test_bozza_si_elimina_con_il_numero(self):
        self.client.force_login(self.compilatore)
        self.assertContains(self.client.get(reverse("analisi", args=[self.analisi.pk])), "Elimina commessa")
        self.client.post(self.url, {"conferma": "sbagliato"})
        self.assertTrue(Commessa.objects.filter(pk=self.commessa.pk).exists())

        risposta = self.client.post(self.url, {"conferma": "26-99"})
        self.assertRedirects(risposta, reverse("elenco_commesse"))
        self.assertFalse(Commessa.objects.filter(numero="26-99").exists())
        self.assertFalse(SchedaAnalisi.objects.exists())
        eliminate = RegistroModifica.objects.filter(azione=RegistroModifica.Azione.ELIMINA)
        self.assertTrue(eliminate.filter(tabella="commessa").exists())
        self.assertFalse(eliminate.filter(tabella="scheda dell'analisi").exists())

    def test_approvata_resta_protetta(self):
        rev = self.analisi.revisione_corrente
        servizi.invia_in_verifica(rev, self.compilatore)
        servizi.segna_verificata(rev, self.verificatore)
        servizi.approva(rev, self.approvatore)
        self.client.force_login(self.compilatore)
        self.assertNotContains(self.client.get(reverse("analisi", args=[self.analisi.pk])), "Elimina commessa")
        self.client.post(self.url, {"conferma": "26-99"})
        self.assertTrue(Commessa.objects.filter(pk=self.commessa.pk).exists())
        with self.assertRaises(ValidationError):
            servizi.elimina_commessa(self.commessa, self.compilatore)

    def test_serve_il_compilatore(self):
        with self.assertRaises(PermissionDenied):
            servizi.elimina_commessa(self.commessa, self.verificatore)
        self.assertTrue(Commessa.objects.filter(pk=self.commessa.pk).exists())


class SoggettiEspostiTest(TestCase):
    """Soggetti esposti (RESS 1.1.1 c e d): libreria, macchina, schede, controlli e documento."""

    @classmethod
    def setUpTestData(cls):
        call_command("importa_libreria", str(LIBRERIA.parent / "Libreria_nuova_Cosmap.xlsx"), stdout=io.StringIO())
        cls.compilatore = crea_utente("mario", servizi.COMPILATORE)
        commessa = Commessa.objects.create(
            numero="26-50", cliente=Cliente.objects.create(ragione_sociale="Prova"), anno=2026,
            riferimento=RiferimentoNormativo.objects.get(),
        )
        cls.macchina = Macchina.objects.create(commessa=commessa, denominazione="Linea", matricola="S1")
        cls.macchina.moduli.set(servizi.moduli_proposti() | Modulo.objects.filter(nome__startswith="Tavola"))
        cls.analisi = servizi.crea_analisi_da_libreria(cls.macchina, cls.compilatore)

    def test_libreria_precompila_i_soggetti(self):
        from .models import Figura

        self.assertEqual(Figura.objects.count(), 8)
        carico = SchedaModello.objects.get(codice="NL-034")
        self.assertIn("Operatore di conduzione", [f.nome for f in carico.soggetti.all()])
        quadro = SchedaModello.objects.get(codice="NL-070")
        self.assertEqual([f.nome for f in quadro.soggetti.all()], ["Manutentore elettrico"])

    def test_soggetti_copiati_e_figure_della_macchina(self):
        rev = self.analisi.revisione_corrente
        scheda = rev.schede.get(codice="NL-034")
        self.assertIn("Operatore di conduzione", [f.nome for f in scheda.soggetti.all()])
        figure = {f.figura.nome: f.descrizione for f in self.macchina.figure.select_related("figura")}
        self.assertIn("Operatore di conduzione", figure)
        self.assertTrue(figure["Operatore di conduzione"])

    def test_controllo_schede_senza_soggetti(self):
        rev = self.analisi.revisione_corrente
        self.assertFalse([a for a in servizi.controlli(rev) if a.tipo == "soggetti"])
        scheda = rev.schede.get(codice="NL-034")
        scheda.soggetti.clear()
        segnalate = [a.scheda for a in servizi.controlli(rev) if a.tipo == "soggetti"]
        self.assertEqual(segnalate, [scheda])

    def test_descrizione_per_macchina_e_documento(self):
        from . import documenti
        from .models import Figura

        self.client.force_login(self.compilatore)
        url = reverse("macchina", args=[self.macchina.pk])
        figure = list(self.macchina.figure.all())
        dati = {
            "denominazione": "Linea", "matricola": "S1", "tipo": "MACCHINA",
            "moduli": [m.pk for m in self.macchina.moduli.all()],
            "figure-TOTAL_FORMS": len(figure) + 1, "figure-INITIAL_FORMS": len(figure),
            "figure-MIN_NUM_FORMS": 0, "figure-MAX_NUM_FORMS": 1000,
        }
        for i, f in enumerate(figure):
            dati.update({f"figure-{i}-id": f.pk, f"figure-{i}-figura": f.figura_id, f"figure-{i}-descrizione": f.descrizione})
        operatore = next(i for i, f in enumerate(figure) if f.figura.nome == "Operatore di conduzione")
        dati[f"figure-{operatore}-descrizione"] = "Addetto alla conduzione linea e carico bancali"
        terzi = next(i for i, f in enumerate(figure) if f.figura.nome == "Terzi di passaggio")
        dati[f"figure-{terzi}-descrizione"] = "Mulettista che transita nella corsia adiacente"
        risposta = self.client.post(url, dati)
        self.assertEqual(risposta.status_code, 302)
        self.assertEqual(
            self.macchina.figure.get(figura=Figura.objects.get(nome="Terzi di passaggio")).descrizione,
            "Mulettista che transita nella corsia adiacente",
        )

        scheda = self.analisi.revisione_corrente.schede.get(codice="NL-034")
        pagina = self.client.get(reverse("scheda", args=[scheda.pk]))
        self.assertContains(pagina, "Addetto alla conduzione linea e carico bancali")

        testo = DocumentiTest.leggi(documenti.valutazione(self.analisi.revisione_corrente))
        self.assertIn("Soggetti esposti", testo)
        self.assertIn("Mulettista che transita nella corsia adiacente", testo)
        self.assertIn("Persona esposta (RESS 1.1.1 c)", testo)

    def test_completa_soggetti_nelle_bozze(self):
        rev = self.analisi.revisione_corrente
        for scheda in rev.schede.all():
            scheda.soggetti.clear()
        self.macchina.figure.all().delete()
        call_command("completa_soggetti", stdout=io.StringIO())
        self.assertFalse([a for a in servizi.controlli(rev) if a.tipo == "soggetti"])
        self.assertTrue(self.macchina.figure.filter(figura__nome="Operatore di conduzione").exists())
