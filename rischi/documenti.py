"""Generazione dei documenti Word di una revisione.

- Dichiarazione UE di conformità (Regolamento (UE) 2023/1230, Allegato V parte A)
- Valutazione dei rischi
- Elenco dei rischi residui per il manuale

I testi fissi della dichiarazione vanno confrontati con il testo ufficiale del
Regolamento prima dell'uso: l'applicazione struttura il documento, la
responsabilità resta di chi firma.
"""

import io
from itertools import groupby

from django.utils import timezone
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .models import Esito, Fabbricante, Revisione, SchedaAnalisi

REGOLAMENTO = "Regolamento (UE) 2023/1230"

TESTI = {
    "it": {
        "titolo": "Dichiarazione UE di conformità",
        "sottotitolo": "Regolamento (UE) 2023/1230 relativo alle macchine, Allegato V, parte A",
        "numero": "N.",
        "fabbricante": "Fabbricante",
        "piva": "Partita IVA",
        "fascicolo": "Persona autorizzata a costituire la documentazione tecnica",
        "responsabilita": (
            "La presente dichiarazione di conformità è rilasciata sotto la responsabilità esclusiva del fabbricante."
        ),
        "oggetto": "Oggetto della dichiarazione",
        "denominazione": "Denominazione",
        "funzione": "Funzione",
        "modello": "Modello / tipo",
        "matricola": "Numero di serie",
        "anno": "Anno di costruzione",
        "conforme": (
            "L'oggetto della dichiarazione sopra descritto è conforme a tutte le disposizioni pertinenti del "
            "Regolamento (UE) 2023/1230 relativo alle macchine"
        ),
        "conforme_altre": " e alla seguente normativa di armonizzazione dell'Unione:",
        "norme_arm": "Norme armonizzate applicate",
        "norme_altre": "Altre norme e specifiche tecniche applicate",
        "organismo": "Organismo notificato",
        "firmato": "Firmato a nome e per conto di",
        "luogo_data": "Luogo e data",
        "nome_qualifica": "Nome e qualifica",
        "firma": "Firma",
        "bozza": "BOZZA – revisione non approvata, documento non valido",
    },
    "en": {
        "titolo": "EU Declaration of Conformity",
        "sottotitolo": "Regulation (EU) 2023/1230 on machinery, Annex V, Part A",
        "numero": "No.",
        "fabbricante": "Manufacturer",
        "piva": "VAT number",
        "fascicolo": "Person authorised to compile the technical documentation",
        "responsabilita": "This declaration of conformity is issued under the sole responsibility of the manufacturer.",
        "oggetto": "Object of the declaration",
        "denominazione": "Name",
        "funzione": "Function",
        "modello": "Model / type",
        "matricola": "Serial number",
        "anno": "Year of construction",
        "conforme": (
            "The object of the declaration described above is in conformity with all the relevant provisions of "
            "Regulation (EU) 2023/1230 on machinery"
        ),
        "conforme_altre": " and with the following Union harmonisation legislation:",
        "norme_arm": "Harmonised standards applied",
        "norme_altre": "Other standards and technical specifications applied",
        "organismo": "Notified body",
        "firmato": "Signed for and on behalf of",
        "luogo_data": "Place and date",
        "nome_qualifica": "Name and function",
        "firma": "Signature",
        "bozza": "DRAFT – revision not approved, document not valid",
    },
}

ESITI_TESTO = {Esito.OK: "OK", Esito.SUGGERITE: "Misure suggerite", Esito.RICHIESTE: "Misure richieste"}
ESITI_COLORE = {Esito.OK: "E6F4E7", Esito.SUGGERITE: "FFF6D6", Esito.RICHIESTE: "FDE8E6"}


# ---------------------------------------------------------------------------
# Utilità
# ---------------------------------------------------------------------------


def _documento():
    doc = Document()
    sezione = doc.sections[0]
    sezione.page_height, sezione.page_width = Cm(29.7), Cm(21.0)
    for margine in ("left_margin", "right_margin"):
        setattr(sezione, margine, Cm(2.0))
    sezione.top_margin = sezione.bottom_margin = Cm(1.8)
    stile = doc.styles["Normal"]
    stile.font.name = "Calibri"
    stile.font.size = Pt(10.5)
    return doc


def _bozza(doc, revisione, testo):
    if revisione.stato == Revisione.Stato.APPROVATA:
        return
    p = doc.sections[0].header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(testo)
    run.bold = True
    run.font.color.rgb = RGBColor(0xB3, 0x26, 0x1E)


def _sfondo(cella, colore):
    proprieta = cella._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), colore)
    proprieta.append(shd)


def _tabella(doc, intestazioni, righe, larghezze=None):
    tabella = doc.add_table(rows=1, cols=len(intestazioni))
    tabella.style = "Table Grid"
    tabella.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cella, testo in zip(tabella.rows[0].cells, intestazioni):
        cella.text = ""
        cella.paragraphs[0].add_run(testo).bold = True
        _sfondo(cella, "EEF0F3")
    for riga in righe:
        celle = tabella.add_row().cells
        for cella, valore in zip(celle, riga):
            cella.text = "" if valore is None else str(valore)
    if larghezze:
        for riga in tabella.rows:
            for cella, larghezza in zip(riga.cells, larghezze):
                cella.width = Cm(larghezza)
    return tabella


def _coppia(doc, etichetta, valore):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run(f"{etichetta}: ").bold = True
    p.add_run(valore or "–")


def _salva(doc):
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _schede_attive(revisione):
    return list(
        revisione.schede.exclude(decisione=SchedaAnalisi.Decisione.SCARTATA)
        .select_related("modulo", "requisito", "revisione__metodo")
        .prefetch_related("misure__norma", "pericoli", "condizioni", "norme", "soggetti")
    )


def _nome(utente):
    if not utente:
        return "–"
    return utente.get_full_name() or utente.username


def norme_applicate(revisione):
    """Norme citate nelle schede attive: (armonizzate, altre), ordinate per codice."""
    norme = {}
    for scheda in _schede_attive(revisione):
        for norma in scheda.norme.all():
            norme[norma.pk] = norma
        for misura in scheda.misure.all():
            if misura.norma:
                norme[misura.norma.pk] = misura.norma
    ordinate = sorted(norme.values(), key=lambda n: n.codice)
    return [n for n in ordinate if n.armonizzata], [n for n in ordinate if not n.armonizzata]


def _edizione(norma):
    edizione = norma.edizione_vigente or norma.edizione_citata
    return f"{norma.codice}:{edizione}" if edizione else norma.codice


# ---------------------------------------------------------------------------
# Dichiarazione UE di conformità
# ---------------------------------------------------------------------------


def dichiarazione(revisione, lingua="it"):
    t = TESTI[lingua]
    fabbricante = Fabbricante.corrente()
    macchina = revisione.analisi.macchina
    commessa = macchina.commessa
    doc = _documento()
    _bozza(doc, revisione, t["bozza"])

    titolo = doc.add_paragraph()
    titolo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titolo.add_run(t["titolo"])
    run.bold = True
    run.font.size = Pt(16)
    sotto = doc.add_paragraph(t["sottotitolo"])
    sotto.alignment = WD_ALIGN_PARAGRAPH.CENTER
    numero = doc.add_paragraph(f"{t['numero']} {commessa.numero}/{revisione.numero}")
    numero.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading(f"1. {t['fabbricante']}", level=2)
    if fabbricante:
        doc.add_paragraph(fabbricante.ragione_sociale)
        doc.add_paragraph(fabbricante.indirizzo)
        if fabbricante.partita_iva:
            _coppia(doc, t["piva"], fabbricante.partita_iva)
    else:
        doc.add_paragraph("Dati del fabbricante non impostati.")

    doc.add_heading(f"2. {t['fascicolo']}", level=2)
    if fabbricante and fabbricante.persona_fascicolo:
        doc.add_paragraph(fabbricante.persona_fascicolo)
        doc.add_paragraph(fabbricante.indirizzo_persona_fascicolo or fabbricante.indirizzo)
    else:
        doc.add_paragraph("–")

    doc.add_heading(f"3. {t['oggetto']}", level=2)
    _coppia(doc, t["denominazione"], macchina.denominazione)
    _coppia(doc, t["funzione"], macchina.funzione)
    _coppia(doc, t["modello"], macchina.modello)
    _coppia(doc, t["matricola"], macchina.matricola)
    _coppia(doc, t["anno"], str(macchina.anno_costruzione or ""))

    doc.add_heading("4.", level=2)
    doc.add_paragraph(t["responsabilita"])

    legislazioni = list(macchina.altre_legislazioni.all())
    doc.add_heading("5.", level=2)
    doc.add_paragraph(t["conforme"] + (t["conforme_altre"] if legislazioni else "."))
    for legge in legislazioni:
        titolo_legge = legge.titolo_en if lingua == "en" and legge.titolo_en else legge.titolo
        doc.add_paragraph(f"{legge.codice} – {titolo_legge}", style="List Bullet")

    armonizzate, altre = norme_applicate(revisione)
    doc.add_heading(f"6. {t['norme_arm']}", level=2)
    if armonizzate:
        for norma in armonizzate:
            doc.add_paragraph(_edizione(norma), style="List Bullet")
    else:
        doc.add_paragraph("–")
    if altre:
        doc.add_paragraph().add_run(t["norme_altre"]).bold = True
        for norma in altre:
            doc.add_paragraph(_edizione(norma), style="List Bullet")

    if macchina.organismo_notificato.strip():
        doc.add_heading(f"7. {t['organismo']}", level=2)
        doc.add_paragraph(macchina.organismo_notificato)

    doc.add_paragraph()
    doc.add_paragraph(f"{t['firmato']}: {fabbricante.ragione_sociale if fabbricante else ''}")
    firma = doc.add_table(rows=3, cols=2)
    luogo = fabbricante.luogo if fabbricante else ""
    data = f"{timezone.localtime(revisione.approvata_il):%d/%m/%Y}" if revisione.approvata_il else "____________"
    valori = [
        (t["luogo_data"], f"{luogo}, {data}"),
        (t["nome_qualifica"], f"{fabbricante.firmatario}, {fabbricante.qualifica_firmatario}" if fabbricante else ""),
        (t["firma"], "\n\n______________________________"),
    ]
    for riga, (etichetta, valore) in zip(firma.rows, valori):
        riga.cells[0].text = etichetta
        riga.cells[1].text = valore
    return _salva(doc)


# ---------------------------------------------------------------------------
# Valutazione dei rischi
# ---------------------------------------------------------------------------


FATTORI = ("Se", "Fr", "Pr", "Av")


def _valore(descrizioni, fattore, valore):
    if valore is None:
        return "–"
    descrizione = descrizioni.get((fattore, valore))
    return f"{valore} – {descrizione}" if descrizione else str(valore)


def _tabella_stima(doc, scheda, descrizioni):
    righe = [
        [fattore] + [_valore(descrizioni, fattore, getattr(scheda, f"{fattore.lower()}_{quale}")) for quale in ("iniziale", "finale")]
        for fattore in FATTORI
    ]
    righe.append(["Cl"] + [_valore({}, "", getattr(scheda, f"cl_{quale}")) for quale in ("iniziale", "finale")])
    esiti = [getattr(scheda, f"esito_{quale}") for quale in ("iniziale", "finale")]
    righe.append(["Esito"] + [ESITI_TESTO.get(e, "–") for e in esiti])
    tabella = _tabella(doc, ["", "Stima iniziale", "Stima finale"], righe, [2, 7.5, 7.5])
    for cella, esito in zip(tabella.rows[-1].cells[1:], esiti):
        if esito:
            _sfondo(cella, ESITI_COLORE[esito])


def valutazione(revisione):
    macchina = revisione.analisi.macchina
    commessa = macchina.commessa
    metodo = revisione.metodo
    doc = _documento()
    _bozza(doc, revisione, TESTI["it"]["bozza"])

    doc.add_heading("Valutazione dei rischi", level=0)
    _coppia(doc, "Commessa", f"{commessa.numero} – {commessa.cliente}")
    _coppia(doc, "Macchina", f"{macchina.denominazione} {macchina.modello}".strip())
    _coppia(doc, "Matricola", macchina.matricola)
    _coppia(doc, "Riferimento normativo", REGOLAMENTO)
    _coppia(doc, "Revisione", f"{revisione.numero} – {revisione.motivo}")
    doc.add_paragraph()
    _tabella(
        doc,
        ["", "Nome", "Data"],
        [
            ["Compilata", _nome(revisione.compilata_da), f"{revisione.creata_il:%d/%m/%Y}"],
            ["Verificata", _nome(revisione.verificata_da), f"{revisione.verificata_il:%d/%m/%Y}" if revisione.verificata_il else "–"],
            ["Approvata", _nome(revisione.approvata_da), f"{revisione.approvata_il:%d/%m/%Y}" if revisione.approvata_il else "–"],
        ],
        [4, 7, 4],
    )

    doc.add_heading("Soggetti", level=1)
    doc.add_paragraph(
        "Operatori (RESS 1.1.1 d): persone incaricate di installare, far funzionare, regolare, pulire, "
        "riparare o spostare la macchina. Persone esposte (RESS 1.1.1 c): chiunque si trovi interamente "
        "o in parte in una zona pericolosa."
    )
    descrizioni_figure = macchina.descrizioni_figure()
    figure_usate = {f for s in _schede_attive(revisione) for f in s.soggetti.all()}
    figure_usate |= {f.figura for f in macchina.figure.select_related("figura")}
    if figure_usate:
        _tabella(
            doc,
            ["Tipo", "Figura", "Chi è su questa macchina"],
            [
                [f.get_tipo_display(), f.nome, descrizioni_figure.get(f.pk) or f.descrizione]
                for f in sorted(figure_usate, key=lambda f: (f.ordine, f.nome))
            ],
            [5, 4, 8],
        )
    else:
        doc.add_paragraph("Soggetti non ancora indicati.")

    doc.add_heading("Metodo di stima", level=1)
    doc.add_paragraph(
        f"{metodo.versione}. Gravità Se da 1 a 4; classe Cl = Fr + Pr + Av (frequenza di esposizione, "
        "probabilità dell'evento pericoloso, possibilità di evitare il danno). L'esito si legge nella matrice."
    )
    fasce = list(metodo.fasce.all())
    righe = []
    for se in (4, 3, 2, 1):
        righe.append([f"Se {se}"] + [ESITI_TESTO[metodo.esito(se, f.cl_min)] for f in fasce])
    tabella = _tabella(doc, ["Gravità"] + [f"Cl {f.cl_min}-{f.cl_max}" for f in fasce], righe)
    for riga in tabella.rows[1:]:
        for cella in riga.cells[1:]:
            esito = next((k for k, v in ESITI_TESTO.items() if v == cella.text), None)
            if esito:
                _sfondo(cella, ESITI_COLORE[esito])

    descrizioni = metodo.descrizioni()
    if descrizioni:
        doc.add_paragraph("Valori dei fattori:")
        _tabella(
            doc,
            ["Fattore", "Valore", "Descrizione"],
            [[f, v, descrizioni[(f, v)]] for f in FATTORI for v in sorted({v for (ff, v) in descrizioni if ff == f}, reverse=True)],
            [2, 2, 13],
        )

    doc.add_heading("Requisiti non applicabili", level=1)
    non_applicabili = revisione.applicabilita.filter(applicabile=False).select_related("requisito")
    if non_applicabili:
        _tabella(doc, ["Requisito", "Motivazione"], [[f"{a.requisito.codice} {a.requisito.titolo}", a.motivazione] for a in non_applicabili], [6, 11])
    else:
        doc.add_paragraph("Tutti i requisiti sono considerati applicabili.")

    doc.add_heading("Schede di valutazione", level=1)
    schede = _schede_attive(revisione)
    for modulo, gruppo in groupby(schede, key=lambda s: s.modulo):
        doc.add_heading(modulo.nome, level=2)
        for s in gruppo:
            doc.add_heading(f"{s.codice or 'Scheda'} – {s.requisito.codice} {s.requisito.titolo}", level=3)
            _coppia(doc, "Zona", " – ".join(v for v in (s.zona_impianto, s.zona_pericolosa) if v))
            _coppia(doc, "Condizioni operative", ", ".join(c.nome for c in s.condizioni.all()))
            _coppia(doc, "Soggetti esposti", ", ".join(f.nome for f in s.soggetti.all()))
            pericoli = list(s.pericoli.all())
            if pericoli:
                doc.add_paragraph().add_run("Pericoli").bold = True
                for p in pericoli:
                    doc.add_paragraph(f"{p.codice} {p.descrizione}", style="List Bullet")
            if s.ha_stima_iniziale or s.ha_stima_finale:
                _tabella_stima(doc, s, descrizioni)
            misure = list(s.misure.all())
            if misure:
                doc.add_paragraph().add_run("Misure di protezione").bold = True
                for m in misure:
                    testo = m.testo + (f" ({m.norma})" if m.norma else "")
                    doc.add_paragraph(testo, style="List Bullet")
            if s.testo_istruzioni:
                _coppia(doc, "Informazioni per le istruzioni / rischio residuo", s.testo_istruzioni)
            norme = ", ".join(n.codice for n in s.norme.all())
            if norme:
                _coppia(doc, "Norme", norme)

    scartate = revisione.schede.filter(decisione=SchedaAnalisi.Decisione.SCARTATA).select_related("requisito")
    if scartate:
        doc.add_heading("Schede proposte e scartate", level=1)
        _tabella(doc, ["Scheda", "Requisito", "Motivazione"], [[s.codice, f"{s.requisito.codice} {s.requisito.titolo}", s.motivazione] for s in scartate], [2.5, 6, 8.5])
    return _salva(doc)


# ---------------------------------------------------------------------------
# Elenco dei rischi residui
# ---------------------------------------------------------------------------


def rischi_residui(revisione):
    macchina = revisione.analisi.macchina
    doc = _documento()
    _bozza(doc, revisione, TESTI["it"]["bozza"])
    doc.add_heading("Rischi residui e informazioni per le istruzioni", level=0)
    doc.add_paragraph(
        f"Commessa {macchina.commessa.numero} · {macchina.denominazione} {macchina.modello}".strip()
        + f" · matricola {macchina.matricola or '–'} · valutazione dei rischi rev. {revisione.numero}"
    )
    schede = [s for s in _schede_attive(revisione) if s.testo_istruzioni.strip()]
    if not schede:
        doc.add_paragraph("Nessun rischio residuo indicato nelle schede.")
    for modulo, gruppo in groupby(schede, key=lambda s: s.modulo):
        doc.add_heading(modulo.nome, level=1)
        for s in gruppo:
            p = doc.add_paragraph()
            p.add_run(f"{s.requisito.codice} {s.requisito.titolo}").bold = True
            if s.esito_finale and s.esito_finale != Esito.OK:
                p.add_run(f"  [{ESITI_TESTO[s.esito_finale]}]")
            doc.add_paragraph(s.testo_istruzioni)
    return _salva(doc)


GENERATORI = {
    "DICHIARAZIONE": dichiarazione,
    "VALUTAZIONE": valutazione,
    "RESIDUI": rischi_residui,
}


def genera(tipo, revisione, lingua="it"):
    """Contenuto .docx del documento; solo la dichiarazione esiste anche in inglese."""
    if tipo == "DICHIARAZIONE":
        return dichiarazione(revisione, lingua)
    return GENERATORI[tipo](revisione)
